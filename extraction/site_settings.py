from openpyxl.workbook.workbook import Workbook

from extraction.tables import clean_value


SITE_SETTINGS_TABLES = {
    "SiteInformation": {
        "table_name": "SIte_Settings_Site_Information",
        "columns": {
            "ZVM Site Name": "ZvmSiteName",
            "Site Location": "SiteLocation",
            "Contact Name": "ContactName",
            "Contact Email": "ContactEmail",
            "Contact Phone": "ContactPhone",
        },
    },
    "Throttling": {
        "table_name": "Site_Settings_Throttling",
        "columns": {
            "ZVM Site Name": "ZvmSiteName",
            "Limited (MB/s)": "LimitedMbps",
            "Time-Based Limited (MB/s)": "TimeBasedLimitedMbps",
            "From": "From",
            "To": "To",
        },
    },
    "Policies": {
        "table_name": "Site_Settings_Policies",
        "columns": {
            "ZVM Site Name": "ZvmSiteName",
            "Failover/Move Commit Policy": "FailoverMoveCommitPolicy",
            "Default Timeout (MInutes)": "DefaultTimeoutMinutes",
            "Default Script Execution Timeout (Sec)": (
                "DefaultScriptExecutionTimeoutSeconds"
            ),
            "Replication Pause Time (Minutes)": "ReplicationPauseTimeMinutes",
            "Enable Replication to Self": "EnableReplicationToSelf",
            "Copy BIOS UUID": "CopyBiosUuid",
            "Copy Instance UUID": "CopyInstanceUuid",
            "Copy vSphere tags": "CopyVsphereTags",
            "Allow SDRS for Recovery VRAs": "AllowSdrsForRecoveryVras",
        },
    },
    "AdvancedResilience": {
        "table_name": "Site_Settings_Advanced_Resilience",
        "columns": {
            "ZVM Site Name": "ZvmSiteName",
            "Enable Offline Recovery Mode": "EnableOfflineRecoveryMode",
        },
    },
    "EncryptionDetection": {
        "table_name": "Site_Settings_Encryption_Detection",
        "columns": {
            "ZVM Site Name": "ZvmSiteName",
            "Encryption Analyzer": "EncryptionAnalyzer",
        },
    },
}


def extract_site_settings(workbook: Workbook) -> dict[str, list[dict]]:
    worksheet = workbook["Site Settings"]

    return {
        section_name: extract_table_rows(
            worksheet,
            settings["table_name"],
            settings["columns"],
        )
        for section_name, settings in SITE_SETTINGS_TABLES.items()
    }


def extract_table_rows(worksheet, table_name: str, column_names: dict) -> list[dict]:
    if table_name not in worksheet.tables:
        raise ValueError(f"Could not find Site Settings table '{table_name}'")

    table = worksheet.tables[table_name]
    cells = worksheet[table.ref]
    headers = [clean_value(cell.value) for cell in cells[0]]
    rows = []

    for excel_row in cells[1:]:
        values_by_header = {
            header: clean_value(cell.value)
            for header, cell in zip(headers, excel_row)
            if header is not None
        }
        output_row = {
            output_name: values_by_header.get(source_name)
            for source_name, output_name in column_names.items()
        }

        if any(value is not None for value in output_row.values()):
            rows.append(output_row)

    return rows
